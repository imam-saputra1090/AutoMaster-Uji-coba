import docx
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

proposal_path = r"C:\Users\MyBook Hype\Downloads\[Template] Proposal Sayembara Pembuatan Bahan Ajar Digital Jenjang SMK.docx"

doc = docx.Document(proposal_path)

print("=== PARAGRAPHS ===")
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"P{i}: {p.text.strip()}")

print("\n=== TABLES ===")
for i, t in enumerate(doc.tables):
    print(f"\nT{i+1}:")
    for r_idx, row in enumerate(t.rows):
        cells_text = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
        print(f"  R{r_idx}: {cells_text}")
