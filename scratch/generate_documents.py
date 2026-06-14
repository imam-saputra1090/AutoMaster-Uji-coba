import docx
import os
import copy
from docx.shared import Inches

def fill_proposal():
    proposal_template = r"C:\Users\MyBook Hype\Downloads\[Template] Proposal Sayembara Pembuatan Bahan Ajar Digital Jenjang SMK.docx"
    proposal_output = r"C:\Users\MyBook Hype\Downloads\Proposal Sayembara Pembuatan Bahan Ajar Digital - AutoMaster.docx"
    proposal_output_workspace = r"c:\Users\MyBook Hype\OneDrive\Documents\AGEN INFODESK 234\GIM\Proposal Sayembara Pembuatan Bahan Ajar Digital - AutoMaster.docx"
    
    if not os.path.exists(proposal_template):
        print(f"Error: Template not found at {proposal_template}")
        return

    doc = docx.Document(proposal_template)
    table = doc.tables[0]
    
    # 1. Fill Identitas Bahan Ajar (Row 1)
    identitas_text = (
        "Judul Bahan Ajar: AutoMaster — Bengkel Virtual: Game Edukasi Interaktif Dasar Otomotif dan Pemeliharaan Sistem Kendaraan Ringan\n"
        "Jenis Bahan Ajar: Gim Edukasi / Simulasi Digital Interaktif (HTML5 PWA)\n"
        "Program Keahlian: Teknik Otomotif / Teknik Kendaraan Ringan (TKR) — Kurikulum Merdeka (Dasar-Dasar Teknik Otomotif, Elemen 1: K3 & Budaya Industri, Elemen 4: Komponen Otomotif, Elemen 6: Alat Ukur Otomotif)\n"
        "Cakupan Materi: Keselamatan dan Kesehatan Kerja (K3) & Budaya Kerja 5R, Klasifikasi dan Penggunaan Peralatan Tangan (Hand Tools & Power Tools), Penggunaan dan Pembacaan Alat Ukur Otomotif (Jangka Sorong, Mikrometer, Dial Indicator, Multimeter), Pemeliharaan Sistem Rem Kendaraan Ringan, Siklus Kerja Motor 4-Langkah, Kelistrikan Dasar Kendaraan.\n"
        "Sasaran Pengguna: Siswa Kelas X SMK (Fase E) dan Guru Pengampu Mata Pelajaran Vokasi Otomotif."
    )
    table.rows[1].cells[0].text = identitas_text
    
    # 2. Fill Deskripsi Umum (Row 3)
    deskripsi_text = (
        "AutoMaster adalah bahan ajar digital berbentuk gim edukasi dan simulasi praktikum virtual berbasis web (HTML5 PWA) yang dirancang untuk siswa Kelas X SMK Teknik Kendaraan Ringan (TKR) sesuai Kurikulum Merdeka Fase E. Gim ini memadukan visualisasi canggih bertema Cyber HUD (Cyberpunk Head-Up Display) dengan pembelajaran teknik otomotif yang komprehensif.\n\n"
        "Gim ini dirancang sebagai solusi atas keterbatasan alat praktik bengkel fisik dan pengganti modul interaktif Flash (.swf) warisan Kemendikbud yang tidak dapat lagi diputar di peramban modern. Fitur-fitur unggulannya meliputi:\n"
        "1. Integrasi WASM Ruffle: Menghidupkan kembali 177 modul animasi otomotif interaktif warisan Kemendikbud secara langsung di browser tanpa instalasi plugin tambahan.\n"
        "2. Simulasi Praktikum Drag & Drop: Siswa melakukan perakitan komponen virtual (seperti rem cakram, tromol, jangka sorong, micrometer) dengan umpan balik visual dan audio instan.\n"
        "3. Kuis Terproteksi (Kiosk/Exam Mode): Fitur kuis anti-curang yang memaksa mode layar penuh, mendeteksi jika siswa berpindah tab/mencari jawaban, mengurangi skor 10% per pelanggaran, dan menutup kuis otomatis pada pelanggaran ketiga.\n"
        "4. Integrasi Database LAN Lokal & Cloud: Sinkronisasi data progres dan nilai siswa secara real-time ke database SQLite lokal laboratorium (untuk penggunaan offline) serta Google Sheets milik guru.\n"
        "5. Gamifikasi Premium: Menggunakan sistem poin XP, lencana pencapaian (Achievements), papan skor kelas (Urutan Prestasi), dan kustomisasi Avatar Mekanik."
    )
    table.rows[3].cells[0].text = deskripsi_text
    
    # 3. Fill Flowchart (Row 5) - Clear and insert flowchart image
    row5_cell = table.rows[5].cells[0]
    row5_cell.text = ""
    p = row5_cell.paragraphs[0]
    p.text = "Berikut adalah alur interaksi dan navigasi antarmuka dalam gim AutoMaster:"
    flowchart_path = r"c:\Users\MyBook Hype\OneDrive\Documents\AGEN INFODESK 234\GIM\assets\flowchart_interaksi.png"
    if os.path.exists(flowchart_path):
        p.add_run().add_picture(flowchart_path, width=Inches(6.0))
    else:
        p.add_run().text = "\n[Gambar Flowchart tidak ditemukan di assets]"
        
    # 4. Fill Rencana Implementasi & Jadwal (Row 7)
    rencana_text = (
        "Strategi Penggunaan Bahan Ajar di Kelas:\n"
        "1. Model Flipped Classroom: Guru menugaskan siswa memainkan Fase Belajar (membaca materi dan menonton animasi SWF Ruffle) dan Fase Simulasi di rumah satu hari sebelum praktik fisik. Hal ini memastikan siswa sudah memahami nama, bentuk, dan fungsi komponen terlebih dahulu.\n"
        "2. Praktik Mandiri Terbimbing di Lab: Server lokal di laboratorium komputer diaktifkan via JALANKAN_GAME.bat. Sebanyak 50 PC klien siswa terhubung ke server LAN. Siswa merakit komponen secara virtual, dan guru memantau keaktifan siswa secara real-time.\n"
        "3. Evaluasi Terproteksi: Guru menyelenggarakan kuis serempak di lab dengan mengaktifkan Kiosk Mode untuk menjamin objektivitas hasil ujian siswa.\n"
        "4. Umpan Balik Dasbor Guru: Guru melihat hasil evaluasi pada Dasbor Pemantauan (/dashboard) untuk menyaring siswa yang butuh remidiasi sebelum diizinkan praktik dengan peralatan bengkel fisik asli.\n\n"
        "Estimasi Waktu dan Jadwal Tahapan Pembuatan (8 Minggu):\n"
    )
    table.rows[7].cells[0].text = rencana_text
    
    # Append the Schedule Table inside Row 7 cell
    r7_cell = table.rows[7].cells[0]
    p_table = r7_cell.add_paragraph()
    p_table.text = "\nTabel Jadwal Rencana Implementasi Pembuatan Gim:"
    
    schedule_table = r7_cell.add_table(rows=9, cols=10)
    try:
        schedule_table.style = 'Table Grid'
    except Exception:
        try:
            schedule_table.style = 'Normal Table'
        except Exception:
            pass
    
    headers = ["No.", "Kegiatan Pembuatan & Uji Coba", "M1", "M2", "M3", "M4", "M5", "M6", "M7", "M8"]
    for col_idx, text in enumerate(headers):
        schedule_table.cell(0, col_idx).text = text
        
    rows_data = [
        ("1", "Analisis Kurikulum & Pengumpulan Aset Otomotif", "X", "X", "", "", "", "", "", ""),
        ("2", "Penyusunan Konten & Restorasi SWF Ruffle WASM", "", "X", "X", "", "", "", "", ""),
        ("3", "Desain UI/UX Cyber HUD & Animasi Ikon SVG", "", "", "X", "X", "", "", "", ""),
        ("4", "Pengembangan Engine Simulasi & Kuis (Kiosk Mode)", "", "", "", "X", "X", "", "", ""),
        ("5", "Integrasi Server LAN & Google Sheets Sync API", "", "", "", "", "X", "X", "", ""),
        ("6", "Pembuatan Halaman Dasbor Pemantauan Guru", "", "", "", "", "", "X", "X", ""),
        ("7", "Uji Coba Lapangan 50 Klien Lab Komputer TKR", "", "", "", "", "", "", "X", "X"),
        ("8", "Evaluasi Kegunaan (SUS) & Penyusunan Laporan", "", "", "", "", "", "", "", "X")
    ]
    
    for row_idx, row_data in enumerate(rows_data):
        for col_idx, text in enumerate(row_data):
            schedule_table.cell(row_idx + 1, col_idx).text = text
            
    # Set custom widths for schedule table columns for better formatting
    for row in schedule_table.rows:
        row.cells[0].width = Inches(0.5)
        row.cells[1].width = Inches(3.5)
        for col_idx in range(2, 10):
            row.cells[col_idx].width = Inches(0.35)

    # Save documents
    doc.save(proposal_output)
    doc.save(proposal_output_workspace)
    print("Filled Proposal saved successfully.")

def fill_storyboard():
    storyboard_template = r"C:\Users\MyBook Hype\Downloads\[Template] Storyboard Sayembara Pembuatan Bahan Ajar Digital Jenjang SMK.docx"
    storyboard_output = r"C:\Users\MyBook Hype\Downloads\Storyboard Sayembara Pembuatan Bahan Ajar Digital - AutoMaster.docx"
    storyboard_output_workspace = r"c:\Users\MyBook Hype\OneDrive\Documents\AGEN INFODESK 234\GIM\Storyboard Sayembara Pembuatan Bahan Ajar Digital - AutoMaster.docx"
    
    if not os.path.exists(storyboard_template):
        print(f"Error: Template not found at {storyboard_template}")
        return

    doc = docx.Document(storyboard_template)
    
    # Store reference to Table 1 (the template structure)
    template_table = doc.tables[0]
    
    # Define data for the 5 scenes
    scenes_data = [
        # Scene 1
        {
            "scene": "Scene 01 — Landing Page / Panduan Awal",
            "treatment": "Layar pembuka game. Pengguna melihat logo game berputar, membaca deskripsi singkat, dan menekan tombol 'Ayo Main Sekarang!' untuk masuk ke login/registrasi, atau klik tombol 'Dasbor Guru' dengan verifikasi password default 'user123'.",
            "visual": (
                "1. Background gelap bermotif grid hologram biru neon.\n"
                "2. Logo AutoMaster berupa roda gigi bersayap yang berputar perlahan.\n"
                "3. Judul 'AutoMaster: Bengkel Otomotif Virtual' dengan font tebal bercahaya.\n"
                "4. Tombol oranye glowing 'Ayo Main Sekarang! 🎮' dan tombol sekunder 'Portal & Laporan Guru 🧑‍🏫'."
            ),
            "narasi": "Selamat Datang di AutoMaster! Kuasai teori dan praktik dasar otomotif di bengkel virtual. Gunakan PC/HP Anda untuk memulai.",
            "sfx": "click_neon.wav (klik tombol), error.wav (jika password salah)",
            "musik": "techno_synth_ambient.mp3 (volume 40%)",
            "ambience": "workshop_ambience.wav (dengung mesin ringan)",
            "normal": "Tombol Ayo Main: warna gradasi oranye-kuning, border neon tipis",
            "hover": "Tombol membesar 5% dan intensitas cahayanya bertambah (neon glow)",
            "hit": "Klik 'Ayo Main' mengalihkan layar ke Login Siswa. Klik 'Portal Guru' menampilkan prompt password.",
            "swipe": "N/A",
            "show": "Layar fade-in selama 0.5 detik saat pertama dimuat.",
            "drag": "N/A"
        },
        # Scene 2
        {
            "scene": "Scene 02 — Registrasi & Login Mandiri Siswa",
            "treatment": "Form pendaftaran atau login bagi siswa menggunakan NIS. Bidang password secara default terisi 'user123' untuk mempercepat pendaftaran massal siswa di laboratorium.",
            "visual": (
                "1. Kotak form glassmorphic semitransparan dengan border neon cyan tipis.\n"
                "2. Bidang input teks (Nama, NIS, WA), dropdown pilihan kelas (X TKR 1, 2, 3), dan input password (terisi bintang/dots).\n"
                "3. Tombol 'Daftar 🚀' dan tautan 'Masuk sebagai Guru'."
            ),
            "narasi": "Daftarkan Akun Siswa Baru. Cukup masukkan NIS, Nama, dan Kelas Anda untuk mulai mengumpulkan poin prestasi!",
            "sfx": "keypress.wav (ketikan keyboard), login_success.wav (sukses login)",
            "musik": "techno_synth_ambient.mp3 (volume diturunkan menjadi 20%)",
            "ambience": "Hening / dengung background tipis",
            "normal": "Input box menyala cyan tipis saat aktif, tombol Daftar: gradasi oranye",
            "hover": "Kursor berubah menjadi pointer pada bidang input dan tombol Daftar",
            "hit": "Klik 'Daftar 🚀' memvalidasi isian dan mengirim data ke database server LAN/SQLite, kemudian membuka Dashboard Utama.",
            "swipe": "N/A",
            "show": "Transisi slide-up halus saat kotak login muncul.",
            "drag": "N/A"
        },
        # Scene 3
        {
            "scene": "Scene 03 — Beranda / Menu Utama (Cyber HUD)",
            "treatment": "Menu navigasi utama siswa setelah login. Menampilkan panel profil (Avatar berputar dengan scanline laser, progress bar XP model power-cell, bintang total) dan grid 8 kartu menu praktikum. Bagian tengah berisi banner yang berganti otomatis (tujuan pembelajaran 1-4).",
            "visual": (
                "1. Latar belakang dot-matrix hologram biru gelap.\n"
                "2. Panel profil HUD di sudut kiri atas. Avatar berputar dengan garis laser vertikal naik-turun.\n"
                "3. Di tengah, banner besar bergambar neon (engine, kunci pas, micrometer) yang meluncur/cross-fade otomatis.\n"
                "4. 8 kartu menu dengan ikon SVG (Main Sekarang, Perpustakaan, Papan Skor, Pengaturan, Kunci Bengkel, Budaya K3, Alat Ukur, Urutan Prestasi) dengan animasi hover unik."
            ),
            "narasi": "Beranda Utama. Selamat belajar, Mekanik! Pilih modul 'Main Sekarang' untuk memulai praktikum atau 'Urutan Prestasi' untuk melihat peringkat kelas.",
            "sfx": "beep_hud.wav (hover kartu menu), slide_whoosh.wav (slide banner berpindah)",
            "musik": "high_tech_synthwave.mp3 (volume 30%)",
            "ambience": "workshop_ambience.wav",
            "normal": "Kartu menu: kaca gelap transparan dengan border neon tipis sesuai warna tema kartu",
            "hover": "Kartu menu memicu sapuan laser scanline vertikal, braket sudut neon menyala, dan ikon SVG teranimasi (roda gigi berputar, kunci pas bergoyang).",
            "hit": "Klik kartu menu membuka modul terkait (misal: 'Main Sekarang' membuka peta level).",
            "swipe": "N/A",
            "show": "Banner berganti slide otomatis setiap 6.0 detik, indikator dots di bawah banner menyala bergantian.",
            "drag": "N/A"
        },
        # Scene 4
        {
            "scene": "Scene 04 — Layar Urutan Prestasi Kelas",
            "treatment": "Halaman peringkat keaktifan siswa. Data diambil dari server guru lokal. Peringkat ditampilkan dalam bentuk grafik batang horizontal (Chart.js) untuk menjaga privasi dengan menyembunyikan nilai numerik (XP/skor).",
            "visual": (
                "1. Panel grafik batang horizontal. Sumbu Y berisi nama-nama siswa, sumbu X adalah visualisasi panjang bar (makin panjang makin tinggi peringkatnya).\n"
                "2. Border panel menyala kuning neon.\n"
                "3. Tombol 'Kembali ke Beranda' di sudut kiri bawah."
            ),
            "narasi": "Papan Urutan Prestasi Kelas. Peringkat dihitung berdasarkan akumulasi XP dan keaktifan Anda di laboratorium otomotif.",
            "sfx": "click_back.wav (klik tombol), draw_chart.wav (grafik tergambar dengan derau digital)",
            "musik": "high_tech_synthwave.mp3 (volume 30%)",
            "ambience": "workshop_ambience.wav",
            "normal": "Grafik batang terisi warna gradien bersinar, tombol Kembali: border ungu neon",
            "hover": "Hover pada batang grafik menampilkan tooltip nama siswa (tanpa menampilkan angka skor/XP), tombol Kembali membesar 3%",
            "hit": "Klik tombol 'Kembali' memicu transisi layar kembali ke Beranda Utama.",
            "swipe": "N/A",
            "show": "Grafik menggambar batangnya secara dinamis (live animation) saat layar dimuat.",
            "drag": "N/A"
        },
        # Scene 5
        {
            "scene": "Scene 05 — Fase Simulasi Praktikum Jangka Sorong",
            "treatment": "Modul praktik interaktif drag & drop. Siswa merakit bagian-bagian jangka sorong (rahang tetap, rahang geser, pengunci, depth probe) ke bayangan target pada benda kerja logam (piston).",
            "visual": (
                "1. Area kerja di sisi kanan menampilkan gambar 3D siluet piston dan jangka sorong kosong bergaris putus-putus (drop zones).\n"
                "2. Di sisi kiri, panel 'Toolbox' berisi komponen jangka sorong.\n"
                "3. Bagian bawah layar terdapat bar info oranye: '0 / 4 komponen terpasang'."
            ),
            "narasi": "Fase Simulasi: Rakitlah Jangka Sorong! Seret komponen rahang tetap, rahang geser, dan skala vernier ke posisi yang tepat pada gambar benda kerja.",
            "sfx": "drag_pickup.wav (drag start), correct_snap.wav (drop sukses), wrong_bounce.wav (drop salah membal kembali), level_complete.wav (simulasi selesai)",
            "musik": "concentration_ambient.mp3 (volume 25%)",
            "ambience": "workshop_ambience.wav",
            "normal": "Kotak komponen di toolbox: border tipis putih. Drop zone target: siluet abu-abu putus-putus",
            "hover": "Drop zone menyala hijau neon saat kartu komponen didekatkan (drag-over highlight)",
            "hit": "N/A",
            "swipe": "N/A",
            "show": "Bar kemajuan di bagian bawah bertambah (misal '1/4', '2/4') secara real-time setiap kali komponen ditempatkan secara benar.",
            "drag": "Siswa menyeret komponen dari toolbox ke drop zone target. Jika posisi salah, komponen membal kembali ke posisi semula di toolbox."
        }
    ]
    
    # We will write the first scene data directly to the template's first table
    def fill_table_data(t, data):
        t.cell(2, 0).text = data["scene"]
        t.cell(2, 1).text = data["treatment"]
        t.cell(4, 0).text = data["visual"]
        t.cell(6, 0).text = data["narasi"]
        t.cell(8, 1).text = data["sfx"]
        t.cell(8, 3).text = data["normal"]
        t.cell(9, 1).text = data["musik"]
        t.cell(9, 3).text = data["hover"]
        t.cell(10, 1).text = data["ambience"]
        t.cell(10, 3).text = data["hit"]
        t.cell(11, 3).text = data["swipe"]
        t.cell(12, 3).text = data["show"]
        t.cell(13, 3).text = data["drag"]
        
    fill_table_data(template_table, scenes_data[0])
    
    # For scenes 2 to 5, we copy the template table, append it, and fill the data
    for i in range(1, 5):
        tbl_xml = copy.deepcopy(template_table._tbl)
        # Add a paragraph spacer
        doc.add_paragraph()
        doc.add_paragraph().text = f"SCENE STORYBOARD BERIKUTNYA:"
        
        # Append the copied table element to body
        doc.element.body.append(tbl_xml)
        new_table = doc.tables[-1]
        
        # Fill data
        fill_table_data(new_table, scenes_data[i])
        
    # Save documents
    doc.save(storyboard_output)
    doc.save(storyboard_output_workspace)
    print("Filled Storyboard saved successfully.")

if __name__ == "__main__":
    fill_proposal()
    fill_storyboard()
